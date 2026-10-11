"""导出层：Excel 多 Sheet 与 CSV 多文件."""

from __future__ import annotations

from openpyxl import load_workbook

from docflow.export import export_csv, export_excel
from docflow.pipeline import run_pipeline


def test_export_excel_contains_report_sheets(settings, purchase_order, order_csv, tmp_path):
    result = run_pipeline(purchase_order, settings, [order_csv])
    target = tmp_path / "out.xlsx"

    export_excel(result, target, purchase_order)

    assert target.exists()
    workbook = load_workbook(target)
    assert {"汇总", "记录", "问题清单", "字段映射"} <= set(workbook.sheetnames)

    header = [cell.value for cell in workbook["记录"][1]]
    assert "order_no" in header
    assert "复核状态" in header


def test_export_csv_writes_one_file_per_section(settings, purchase_order, order_csv, tmp_path):
    result = run_pipeline(purchase_order, settings, [order_csv])

    files = export_csv(result, tmp_path / "csv", purchase_order)

    assert {path.name for path in files} == {
        "summary.csv",
        "records.csv",
        "issues.csv",
        "mapping.csv",
    }
    assert all(path.exists() for path in files)


def test_export_row_number_stays_integer(settings, purchase_order, tmp_path):
    # 同时制造「行级问题」和「表级问题（整列为空）」：表级问题没有行号，
    # 混入 None 曾把整列升成 float，导出成「1.0」，与 records.csv 的「1」不一致
    path = tmp_path / "mixed.csv"
    path.write_text(
        "订单编号,供应商,物料编号,物料名称,数量,单价,总金额,下单日期\n"
        "PO-2024-001,华信机械,A001,螺栓,10,1,,2024-03-01\n",
        encoding="utf-8",
    )
    result = run_pipeline(purchase_order, settings, [path])

    export_csv(result, tmp_path / "csv", purchase_order)

    issues = (tmp_path / "csv" / "issues.csv").read_text(encoding="utf-8-sig")
    assert "1.0" not in issues
