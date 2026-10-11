"""去重：精确主键分组 + 模糊相似分组（模糊必须人工确认，绝不自动合并）."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from docflow.export import export_excel
from docflow.models.dedupe import DuplicateMode
from docflow.models.record import IssueCode, ReviewStatus
from docflow.pipeline import run_pipeline
from docflow.store import build_session_factory, load_run_result, save_run

HEADER = "订单编号,供应商,物料编号,物料名称,数量,单价,总金额,下单日期,交付日期,是否加急,备注\n"


def _relaxed_primary_key(schema):
    """把主键字段改为非必填，隔离出「去重本身导致复核」这一行为."""
    return schema.model_copy(
        update={
            "fields": [
                field.model_copy(update={"required": False}) if field.key == schema.primary_key else field
                for field in schema.fields
            ]
        }
    )


def test_exact_duplicate_group_by_primary_key(settings, purchase_order, write_csv):
    path = write_csv(
        "dup.csv",
        HEADER
        + "PO-2024-001,华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,\n"
        + "PO-2024-001,远大五金,A002,支架,5,2,10,2024-03-02,2024-03-20,否,\n",
    )

    result = run_pipeline(purchase_order, settings, [path])

    assert len(result.dedupe.exact_groups) == 1
    group = result.dedupe.exact_groups[0]
    assert group.mode is DuplicateMode.EXACT
    assert group.size == 2
    assert not group.review_required
    assert result.duplicate_records == 1


def test_fuzzy_duplicate_forces_review(settings, purchase_order, write_csv):
    schema = _relaxed_primary_key(purchase_order)
    path = write_csv(
        "fuzzy.csv",
        HEADER
        + ",华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,第一条\n"
        + ",华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,第二条\n",
    )

    result = run_pipeline(schema, settings, [path])

    assert len(result.dedupe.fuzzy_groups) == 1
    group = result.dedupe.fuzzy_groups[0]
    assert group.review_required
    assert group.similarity >= 0.9
    assert result.dedupe.review_required_count == 2
    # 去重前这两条本可自动通过（置信度 1.0），是去重把它们拉回了复核队列
    for record in result.records:
        assert record.review_status is ReviewStatus.PENDING
        assert any(issue.code is IssueCode.POSSIBLE_DUPLICATE for issue in record.issues)


def test_reliable_primary_key_skips_fuzzy_matching(settings, purchase_order, write_csv):
    """主键可靠时不做模糊匹配，避免把不同订单误判为重复."""
    path = write_csv(
        "ok.csv",
        HEADER
        + "PO-2024-001,华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,\n"
        + "PO-2024-002,华信机械,A001,螺栓,100,0.5,50,2024-03-02,2024-03-16,否,\n",
    )

    result = run_pipeline(purchase_order, settings, [path])

    assert not result.dedupe.groups


def test_dedupe_disabled_by_schema(settings, purchase_order, write_csv):
    off = purchase_order.model_copy(update={"dedupe": purchase_order.dedupe.model_copy(update={"enabled": False})})
    path = write_csv(
        "dup.csv",
        HEADER
        + "PO-2024-001,华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,\n"
        + "PO-2024-001,远大五金,A002,支架,5,2,10,2024-03-02,2024-03-20,否,\n",
    )

    result = run_pipeline(off, settings, [path])

    assert not result.dedupe.groups


def test_dedupe_group_is_exported(settings, purchase_order, write_csv, tmp_path: Path):
    path = write_csv(
        "dup.csv",
        HEADER
        + "PO-2024-001,华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,\n"
        + "PO-2024-001,远大五金,A002,支架,5,2,10,2024-03-02,2024-03-20,否,\n",
    )
    result = run_pipeline(purchase_order, settings, [path])

    target = tmp_path / "out.xlsx"
    export_excel(result, target, purchase_order)

    workbook = load_workbook(target)
    assert "去重组" in workbook.sheetnames
    header = [cell.value for cell in workbook["去重组"][1]]
    assert "实体标识" in header


def test_dedupe_groups_survive_persistence(settings, purchase_order, write_csv):
    path = write_csv(
        "dup.csv",
        HEADER
        + "PO-2024-001,华信机械,A001,螺栓,100,0.5,50,2024-03-01,2024-03-15,否,\n"
        + "PO-2024-001,远大五金,A002,支架,5,2,10,2024-03-02,2024-03-20,否,\n",
    )
    result = run_pipeline(purchase_order, settings, [path])

    factory = build_session_factory(settings)
    with factory() as session:
        save_run(session, result)
        session.commit()
    with factory() as session:
        loaded = load_run_result(session, result.run_id)

    assert loaded is not None
    assert len(loaded.dedupe.groups) == 1
    assert loaded.dedupe.groups[0].mode is DuplicateMode.EXACT
